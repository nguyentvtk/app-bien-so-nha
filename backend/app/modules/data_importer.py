"""
Module Nhập Liệu Từ File Excel Thực Tế — UBND XÃ TÂN PHÚ, TỈNH TÂY NINH
Trích xuất, chuẩn hóa và nạp dữ liệu từ 04 tệp Excel số liệu thực địa:
- DS_Nha_thuc_te.xlsx: 7.400 hộ dân trên 11 ấp (Bảng kê hộ điều tra)
- Trang_1.xlsx, Trang_2.xlsx, Trang_3.xlsx: 6.491 thửa đất địa chính (KH 2959/KH-BNNMT-BCA)
"""

import os
import re
import json
import random
import logging
import openpyxl
from typing import List, Dict, Any, Optional, Tuple

logger = logging.getLogger(__name__)

# Đường dẫn mặc định tới thư mục chứa file Excel trên Google Drive
DEFAULT_EXCEL_FOLDER = os.path.join(
    os.path.expanduser("~"),
    "Library/CloudStorage/GoogleDrive-nguyentvtk@gmail.com/Drive của tôi/App/AppCapBienSoNha"
)

# Mã loại đất có chứa đất ở (ONT - Đất ở nông thôn)
MA_LOAI_DAT_CO_THO_CU = {
    "ONT", "ONT+CLN", "CLN+ONT", "ONT+HNK", "HNK+ONT",
    "ONT+LUK", "LUK+ONT", "ONT+BHK", "BHK+ONT",
    "ONT+CLN+HNK", "ONT+HNK+CLN", "ONT+CLN+LUK",
    "ONT+NTS", "ONT+DTL", "ONT+DGT"
}

# Mã loại đất nông nghiệp thuần (không có ONT)
MA_LOAI_DAT_NONG_NGHIEP = {
    "HNK", "CLN", "LUK", "BHK", "NTS", "DTL", "DGT", "DGD", "DNL",
    "CLN+HNK", "HNK+CLN", "CLN+LUK", "LUK+CLN", "HNK+LUK",
    "CLN+BHK", "BHK+CLN", "HNK+BHK"
}


def normalize_name(name: Optional[str]) -> str:
    """Chuẩn hóa tên người để đối soát linh hoạt giữa 2 CSDL."""
    if not name:
        return ""
    s = str(name).strip()
    # Xóa tiền tố phổ biến trong CSDL địa chính
    s = re.sub(r'^(Ông|Bà|Hộ ông|Hộ bà|Ông bà|Bà và ông)\s+', '', s, flags=re.IGNORECASE)
    # Lấy người đầu tiên nếu có "và"
    s = re.sub(r'\s+và\s+.*$', '', s, flags=re.IGNORECASE)
    # Chuẩn hóa khoảng trắng
    s = re.sub(r'\s+', ' ', s).strip()
    return s


def normalize_name_lower(name: Optional[str]) -> str:
    """Chuẩn hóa tên và chuyển thành chữ thường để so khớp."""
    return normalize_name(name).lower()


def kiem_tra_co_dat_o(ma_loai_dat: Optional[str]) -> bool:
    """Kiểm tra mã loại đất có chứa thổ cư (ONT) hay không."""
    if not ma_loai_dat:
        return False
    s = str(ma_loai_dat).strip().upper()
    return "ONT" in s


def phan_loai_phap_ly_thua_dat(ma_loai_dat: Optional[str]) -> str:
    """Phân loại pháp lý theo mã loại đất cho mục đích cấp biển số nhà."""
    if not ma_loai_dat:
        return "KHONG_XAC_DINH"
    s = str(ma_loai_dat).strip().upper()
    if "ONT" in s:
        return "DAT_O_THO_CU"
    elif any(m in s for m in ["DGT", "DGD", "DNL", "DTL"]):
        return "DAT_CONG_CONG"
    else:
        return "DAT_NONG_NGHIEP"


class DataImporter:
    """Module đọc và nạp dữ liệu từ 04 file Excel số liệu thực địa vào CSDL SQLite."""

    def __init__(self, storage_service):
        self.storage = storage_service

    # ====================== NHẬP THỬA ĐẤT ĐỊA CHÍNH ======================

    def import_thua_dat(self, folder: str = DEFAULT_EXCEL_FOLDER) -> Dict[str, Any]:
        """
        Nhập 6.491 thửa đất địa chính từ Trang_1.xlsx, Trang_2.xlsx, Trang_3.xlsx.
        Schema cột Excel (header dòng 4-6, dữ liệu từ dòng 7):
          Col 1: STT
          Col 2: Tên đơn vị hành chính (Xã Tân Phú)
          Col 3: Số tờ bản đồ
          Col 4: Số thứ tự thửa đất
          Col 5: Dữ liệu đã được xuất sổ địa chính điện tử (X)
          Col 6: Dữ liệu chưa xuất sổ
          Col 7: Đồng bộ 3 khối (X)
          Col 8: Không đồng bộ 3 khối
          Col 9: Chỉ có dữ liệu thuộc tính
          Col 10: Khớp CSDL QG dân cư
          Col 11: Chưa khớp
          Col 12: Không xác định
          Col 13: Đang vận hành 24/7
          Col 14: Không vận hành
          Col 15: Nhóm phân loại (Nhóm 1, 2, 3)
          Col 16: Thông tin chủ (Họ tên chủ sử dụng đất)
          Col 17: Thông tin GCN (số GCN)
          Col 18: Đã định danh thửa
          Col 19: Đã quy đổi
          Col 20: Thông tin thiếu (chi tiết)
          Col 21: Diện tích (m2)
          Col 22: Mã loại đất (ONT, HNK, CLN+ONT...)
        """
        files = ["Trang_1.xlsx", "Trang_2.xlsx", "Trang_3.xlsx"]
        total = 0
        errors = []

        conn = self.storage.get_connection()
        cursor = conn.cursor()

        # Xóa dữ liệu cũ để nạp lại sạch sẽ
        cursor.execute("DELETE FROM thua_dat")

        for fn in files:
            fp = os.path.join(folder, fn)
            if not os.path.exists(fp):
                errors.append(f"Không tìm thấy tệp {fn}")
                continue

            logger.info(f"Đang đọc {fn}...")
            # Phải dùng read_only=False vì file có merged cells phức tạp
            wb = openpyxl.load_workbook(fp, data_only=True, read_only=False)
            sh = wb[wb.sheetnames[0]]

            for r in range(7, sh.max_row + 1):
                so_to_raw = sh.cell(row=r, column=3).value
                so_thua_raw = sh.cell(row=r, column=4).value

                if so_to_raw is None or so_thua_raw is None:
                    continue

                so_to = str(so_to_raw).strip()
                so_thua = str(so_thua_raw).strip()

                if not so_to or not so_thua:
                    continue

                ten_chu = normalize_name(sh.cell(row=r, column=16).value)
                so_gcn_raw = sh.cell(row=r, column=17).value
                so_gcn = str(so_gcn_raw).strip() if so_gcn_raw else None
                thong_tin_thieu_raw = sh.cell(row=r, column=20).value
                thong_tin_thieu = str(thong_tin_thieu_raw).strip() if thong_tin_thieu_raw else None
                nhom_raw = sh.cell(row=r, column=15).value
                nhom_phan_loai = str(nhom_raw).strip() if nhom_raw else None

                # Diện tích (cột 21)
                dien_tich_raw = sh.cell(row=r, column=21).value
                try:
                    dien_tich = float(str(dien_tich_raw).replace(",", ".")) if dien_tich_raw else 0.0
                except (ValueError, TypeError):
                    dien_tich = 0.0

                # Mã loại đất (cột 22)
                loai_dat_raw = sh.cell(row=r, column=22).value
                ma_loai_dat = str(loai_dat_raw).strip().upper() if loai_dat_raw else ""
                co_dat_o = 1 if kiem_tra_co_dat_o(ma_loai_dat) else 0

                # Các cờ Boolean (X = có)
                da_xuat_so = 1 if str(sh.cell(row=r, column=5).value or "").strip().upper() == "X" else 0
                dong_bo_3k = 1 if str(sh.cell(row=r, column=7).value or "").strip().upper() == "X" else 0
                lam_sach = 1 if str(sh.cell(row=r, column=10).value or "").strip().upper() == "X" else 0
                van_hanh = 1 if str(sh.cell(row=r, column=13).value or "").strip().upper() == "X" else 0
                dinh_danh = 1 if sh.cell(row=r, column=18).value else 0

                try:
                    cursor.execute("""
                    INSERT OR REPLACE INTO thua_dat (
                        so_to, so_thua, ten_chu, ma_loai_dat, dien_tich, so_gcn,
                        nhom_phan_loai, da_xuat_so_dien_tu, dong_bo_3_khoi,
                        lam_sach_csdl_dan_cu, van_hanh_24_7, da_dinh_danh_thua,
                        thong_tin_thieu, co_dat_o, nguon_du_lieu
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """, (
                        so_to, so_thua, ten_chu, ma_loai_dat, dien_tich, so_gcn,
                        nhom_phan_loai, da_xuat_so, dong_bo_3k,
                        lam_sach, van_hanh, dinh_danh,
                        thong_tin_thieu, co_dat_o, f"KH2959_{fn}"
                    ))
                    total += 1
                except Exception as e:
                    errors.append(f"{fn} Tờ {so_to} Thửa {so_thua}: {str(e)}")

            wb.close()

        # Log thao tác
        cursor.execute(
            "INSERT INTO log_hethong (hanh_dong, doi_tuong, chi_tiet) VALUES (?, ?, ?)",
            ("NHẬP THỬA ĐẤT ĐỊA CHÍNH", "thua_dat",
             f"Đã nạp {total} thửa đất từ 3 file KH2959/BCA. Lỗi: {len(errors)}")
        )
        conn.commit()
        conn.close()

        self.storage._cache.clear()
        return {
            "success": True,
            "message": f"Đã nạp thành công {total} thửa đất địa chính xã Tân Phú vào CSDL!",
            "total": total,
            "errors": errors[:10] if errors else []
        }

    # ====================== NHẬP HỘ DÂN THỰC TẾ ======================

    def import_ho_dan(self, folder: str = DEFAULT_EXCEL_FOLDER) -> Dict[str, Any]:
        """
        Nhập 7.400 hộ dân từ DS_Nha_thuc_te.xlsx (11 ấp x 1 sheet mỗi ấp).
        Schema cột Excel (header dòng 8, dữ liệu từ dòng 11):
          Col 1: STT nhà (A)
          Col 2: STT hộ (B)
          Col 3: Họ và tên chủ hộ (C)
          Col 4: Địa chỉ hộ (D)
          Col 5: Số nhân khẩu (1)
          Col 6: Ghi chú (2)
        Thông tin điều tra viên ở dòng 6 của sheet (Người lập bảng kê).
        Tên ấp trích từ tên sheet.
        """
        fp = os.path.join(folder, "DS_Nha_thuc_te.xlsx")
        if not os.path.exists(fp):
            return {"success": False, "message": f"Không tìm thấy tệp DS_Nha_thuc_te.xlsx"}

        wb = openpyxl.load_workbook(fp, data_only=True, read_only=False)
        total = 0
        ap_stats = {}
        errors = []

        conn = self.storage.get_connection()
        cursor = conn.cursor()

        # Xóa dữ liệu cũ
        cursor.execute("DELETE FROM ho_dan")

        for sname in wb.sheetnames:
            # Bỏ qua sheet tổng hợp "Bảng kê - Tân Phú"
            if "bảng kê" in sname.lower() or "bang ke" in sname.lower():
                continue

            sh = wb[sname]

            # Trích xuất tên ấp từ tên sheet (VD: "1,Tân Đông" -> "Tân Đông")
            ap_name = re.sub(r'^\s*\d+[\s,]*', '', sname).strip()
            if not ap_name:
                ap_name = sname.strip()

            # Trích xuất thông tin Điều tra viên (dòng 6)
            dtv_raw = str(sh.cell(row=6, column=1).value or "")
            dtv_info = ""
            dtv_match = re.search(r':\s*(.+)', dtv_raw)
            if dtv_match:
                dtv_info = dtv_match.group(1).strip()

            sdt_raw = str(sh.cell(row=6, column=5).value or "")
            dtv_sdt = ""
            sdt_match = re.search(r':\s*(.+)', sdt_raw)
            if sdt_match:
                dtv_sdt = sdt_match.group(1).strip()

            ap_count = 0
            for r in range(11, sh.max_row + 1):
                ten_chu_ho_raw = sh.cell(row=r, column=3).value
                if not ten_chu_ho_raw:
                    continue

                ten_chu_ho = str(ten_chu_ho_raw).strip()
                if not ten_chu_ho:
                    continue

                stt_nha = None
                stt_ho = None
                try:
                    stt_nha = int(sh.cell(row=r, column=1).value) if sh.cell(row=r, column=1).value else None
                except (ValueError, TypeError):
                    pass
                try:
                    stt_ho = int(sh.cell(row=r, column=2).value) if sh.cell(row=r, column=2).value else None
                except (ValueError, TypeError):
                    pass

                dia_chi_raw = sh.cell(row=r, column=4).value
                dia_chi = str(dia_chi_raw).strip() if dia_chi_raw else ""
                try:
                    so_nhan_khau = int(sh.cell(row=r, column=5).value) if sh.cell(row=r, column=5).value else 0
                except (ValueError, TypeError):
                    so_nhan_khau = 0
                ghi_chu_raw = sh.cell(row=r, column=6).value
                ghi_chu = str(ghi_chu_raw).strip() if ghi_chu_raw else None

                try:
                    cursor.execute("""
                    INSERT INTO ho_dan (
                        stt_nha, stt_ho, ten_chu_ho, dia_chi, ap,
                        so_nhan_khau, ghi_chu, dieu_tra_vien, sdt_dieu_tra_vien,
                        nguon_du_lieu
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """, (
                        stt_nha, stt_ho, ten_chu_ho, dia_chi, ap_name,
                        so_nhan_khau, ghi_chu, dtv_info, dtv_sdt,
                        "DS_Nha_thuc_te.xlsx"
                    ))
                    total += 1
                    ap_count += 1
                except Exception as e:
                    errors.append(f"Ấp {ap_name} - {ten_chu_ho}: {str(e)}")

            ap_stats[ap_name] = ap_count

        wb.close()

        cursor.execute(
            "INSERT INTO log_hethong (hanh_dong, doi_tuong, chi_tiet) VALUES (?, ?, ?)",
            ("NHẬP HỘ DÂN THỰC TẾ", "ho_dan",
             f"Đã nạp {total} hộ dân từ 11 ấp xã Tân Phú. Lỗi: {len(errors)}")
        )
        conn.commit()
        conn.close()

        self.storage._cache.clear()
        return {
            "success": True,
            "message": f"Đã nạp thành công {total} hộ dân trên 11 ấp xã Tân Phú!",
            "total": total,
            "ap_stats": ap_stats,
            "errors": errors[:10] if errors else []
        }

    # ====================== ĐỐI SOÁT CHÉO ======================

    def doi_soat_cheo(self) -> Dict[str, Any]:
        """
        Chạy thuật toán đối soát chéo giữa Hộ dân (7.400) và Thửa đất (6.491):
        - Khớp tên chủ hộ <-> tên chủ sử dụng đất.
        - Phân loại pháp lý: Đủ điều kiện / Chưa đủ điều kiện / Chưa đối soát.
        """
        conn = self.storage.get_connection()
        cursor = conn.cursor()

        # 1. Xây dựng bảng tra cứu tên chủ đất (chuẩn hóa chữ thường)
        cursor.execute("SELECT id, so_to, so_thua, ten_chu, ma_loai_dat, co_dat_o, so_gcn FROM thua_dat")
        land_index = {}
        for row in cursor.fetchall():
            r = dict(row)
            key = normalize_name_lower(r["ten_chu"])
            if key:
                if key not in land_index:
                    land_index[key] = []
                land_index[key].append(r)

        # 2. Đối soát từng hộ dân
        cursor.execute("SELECT id, ten_chu_ho FROM ho_dan")
        all_ho = cursor.fetchall()
        total = len(all_ho)
        matched = 0
        du_dk = 0
        chua_du_dk = 0
        chua_doi_soat = 0

        for ho in all_ho:
            ho_id = ho["id"]
            key = normalize_name_lower(ho["ten_chu_ho"])

            if key in land_index:
                matched += 1
                thua_list = land_index[key]
                # Ưu tiên thửa đất có đất ở (ONT)
                thua_ont = [t for t in thua_list if t["co_dat_o"] == 1]

                if thua_ont:
                    best = thua_ont[0]
                    phan_loai = "DU_DIEU_KIEN"
                    du_dk += 1
                else:
                    best = thua_list[0]
                    phan_loai = "CHUA_DU_DIEU_KIEN"
                    chua_du_dk += 1

                lien_ket = f"Tờ {best['so_to']}, Thửa {best['so_thua']} ({best['ma_loai_dat']})"
                if best.get("so_gcn"):
                    lien_ket += f" | GCN: {best['so_gcn']}"

                cursor.execute(
                    "UPDATE ho_dan SET thua_dat_lien_ket = ?, phan_loai_phap_ly = ? WHERE id = ?",
                    (lien_ket, phan_loai, ho_id)
                )
            else:
                chua_doi_soat += 1
                cursor.execute(
                    "UPDATE ho_dan SET phan_loai_phap_ly = 'CHUA_DOI_SOAT' WHERE id = ?",
                    (ho_id,)
                )

        cursor.execute(
            "INSERT INTO log_hethong (hanh_dong, doi_tuong, chi_tiet) VALUES (?, ?, ?)",
            ("ĐỐI SOÁT CHÉO", "ho_dan + thua_dat",
             f"Khớp: {matched}/{total} ({matched*100//total}%). "
             f"Đủ ĐK: {du_dk}. Chưa đủ ĐK: {chua_du_dk}. Chưa đối soát: {chua_doi_soat}")
        )
        conn.commit()
        conn.close()

        self.storage._cache.clear()
        return {
            "success": True,
            "message": f"Hoàn tất đối soát chéo {total} hộ dân với {len(land_index)} chủ đất!",
            "tong_ho_dan": total,
            "khop_ten": matched,
            "ti_le_khop": f"{matched * 100 / total:.1f}%",
            "du_dieu_kien_cap_so": du_dk,
            "chua_du_dieu_kien": chua_du_dk,
            "chua_doi_soat": chua_doi_soat
        }

    # ====================== THỐNG KÊ ======================

    def thong_ke_thua_dat(self) -> Dict[str, Any]:
        """Thống kê tổng hợp thửa đất địa chính."""
        conn = self.storage.get_connection()
        cursor = conn.cursor()

        cursor.execute("SELECT COUNT(*) FROM thua_dat")
        tong = cursor.fetchone()[0]

        cursor.execute("SELECT COUNT(*) FROM thua_dat WHERE co_dat_o = 1")
        co_dat_o = cursor.fetchone()[0]

        cursor.execute("SELECT COUNT(*) FROM thua_dat WHERE co_dat_o = 0")
        khong_dat_o = cursor.fetchone()[0]

        cursor.execute("SELECT ma_loai_dat, COUNT(*) as cnt FROM thua_dat GROUP BY ma_loai_dat ORDER BY cnt DESC LIMIT 15")
        phan_bo = {row["ma_loai_dat"]: row["cnt"] for row in cursor.fetchall()}

        conn.close()
        return {
            "tong_thua_dat": tong,
            "co_dat_o_ont": co_dat_o,
            "dat_nong_nghiep_thuan": khong_dat_o,
            "phan_bo_loai_dat": phan_bo
        }

    def thong_ke_ho_dan(self) -> Dict[str, Any]:
        """Thống kê tổng hợp hộ dân theo ấp và phân loại pháp lý."""
        conn = self.storage.get_connection()
        cursor = conn.cursor()

        cursor.execute("SELECT COUNT(*) FROM ho_dan")
        tong = cursor.fetchone()[0]

        cursor.execute("SELECT ap, COUNT(*) as cnt FROM ho_dan GROUP BY ap ORDER BY cnt DESC")
        theo_ap = {row["ap"]: row["cnt"] for row in cursor.fetchall()}

        cursor.execute("SELECT phan_loai_phap_ly, COUNT(*) as cnt FROM ho_dan GROUP BY phan_loai_phap_ly")
        theo_pl = {row["phan_loai_phap_ly"]: row["cnt"] for row in cursor.fetchall()}

        cursor.execute("SELECT SUM(so_nhan_khau) FROM ho_dan")
        tong_nhan_khau = cursor.fetchone()[0] or 0

        conn.close()
        return {
            "tong_ho_dan": tong,
            "tong_nhan_khau": tong_nhan_khau,
            "theo_ap": theo_ap,
            "phan_loai_phap_ly": theo_pl
        }

    def search_ho_dan(self, query: str, limit: int = 50) -> List[Dict[str, Any]]:
        """Tìm kiếm hộ dân theo tên, địa chỉ hoặc ấp."""
        conn = self.storage.get_connection()
        cursor = conn.cursor()
        q = f"%{query.strip().lower()}%"
        cursor.execute("""
            SELECT * FROM ho_dan
            WHERE LOWER(ten_chu_ho) LIKE ?
               OR LOWER(COALESCE(dia_chi, '')) LIKE ?
               OR LOWER(ap) LIKE ?
               OR LOWER(COALESCE(thua_dat_lien_ket, '')) LIKE ?
            ORDER BY ap, stt_nha
            LIMIT ?
        """, (q, q, q, q, limit))
        results = [dict(r) for r in cursor.fetchall()]
        conn.close()
        return results

    def search_thua_dat(self, query: str, limit: int = 50) -> List[Dict[str, Any]]:
        """Tìm kiếm thửa đất theo tờ, thửa, tên chủ hoặc mã loại đất."""
        conn = self.storage.get_connection()
        cursor = conn.cursor()
        q = f"%{query.strip().lower()}%"
        cursor.execute("""
            SELECT * FROM thua_dat
            WHERE so_to LIKE ?
               OR so_thua LIKE ?
               OR LOWER(COALESCE(ten_chu, '')) LIKE ?
               OR LOWER(COALESCE(ma_loai_dat, '')) LIKE ?
               OR LOWER(COALESCE(so_gcn, '')) LIKE ?
            ORDER BY CAST(so_to AS INTEGER), CAST(so_thua AS INTEGER)
            LIMIT ?
        """, (q, q, q, q, q, limit))
        results = [dict(r) for r in cursor.fetchall()]
        conn.close()
        return results

    def get_ho_dan_by_ap(self, ap: str) -> List[Dict[str, Any]]:
        """Lấy danh sách hộ dân theo ấp."""
        conn = self.storage.get_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM ho_dan WHERE ap = ? ORDER BY stt_nha, stt_ho", (ap,))
        results = [dict(r) for r in cursor.fetchall()]
        conn.close()
        return results

    def get_danh_sach_ap(self) -> List[Dict[str, Any]]:
        """Trả về danh sách 11 ấp và số hộ mỗi ấp."""
        conn = self.storage.get_connection()
        cursor = conn.cursor()
        cursor.execute("""
            SELECT ap, COUNT(*) as so_ho, SUM(so_nhan_khau) as so_nhan_khau,
                   SUM(CASE WHEN phan_loai_phap_ly = 'DU_DIEU_KIEN' THEN 1 ELSE 0 END) as du_dk,
                   SUM(CASE WHEN phan_loai_phap_ly = 'CHUA_DU_DIEU_KIEN' THEN 1 ELSE 0 END) as chua_du_dk,
                   SUM(CASE WHEN phan_loai_phap_ly = 'CHUA_DOI_SOAT' THEN 1 ELSE 0 END) as chua_doi_soat
            FROM ho_dan GROUP BY ap ORDER BY ap
        """)
        results = [dict(r) for r in cursor.fetchall()]
        conn.close()
        return results

    # ====================== NẠP DỮ LIỆU VÀO SỐ NHÀ ======================

    # Tọa độ trung tâm (centroid WGS84) của mỗi ấp xã Tân Phú, Tây Ninh
    AP_CENTROIDS = {
        "Tân Đông":   {"lat": 11.6120, "lng": 106.1780},
        "Tân Hòa":    {"lat": 11.6010, "lng": 106.1920},
        "Tân Tiến":   {"lat": 11.5980, "lng": 106.1850},
        "Tân Xuân":   {"lat": 11.6080, "lng": 106.1950},
        "Tân Thạnh":  {"lat": 11.5950, "lng": 106.1780},
        "Tân Trung A": {"lat": 11.6050, "lng": 106.1850},
        "Tân Trung B": {"lat": 11.6030, "lng": 106.1870},
        "Tân Lợi A":  {"lat": 11.6100, "lng": 106.1700},
        "Tân Lợi B":  {"lat": 11.6070, "lng": 106.1720},
        "Tân Tây":    {"lat": 11.5920, "lng": 106.1900},
        "Tân Châu":   {"lat": 11.6150, "lng": 106.1680},
    }

    # Mẫu regex để trích xuất tên đường từ địa chỉ hộ dân
    # Khớp: ĐT785, ĐT.785, DT785, ĐH809, DH 819, ĐH.781B, THU35, THU.92, THU94Đ, QL22C...
    ROAD_PATTERN = re.compile(
        r'(ĐT|DT|ĐH|DH|QL|TL)[\.\s]?(\d+[A-ZĐa-zđ]?)',
        re.IGNORECASE
    )
    THU_PATTERN = re.compile(
        r'THU[\.\s]?(\d+[A-ZĐa-zđ]*)',
        re.IGNORECASE
    )

    def _tao_tuyen_duong_ap(self, cursor, ap_name: str) -> str:
        """Tạo tuyến đường mặc định cấp ấp nếu chưa có. Trả về id_tuyen."""
        # Chuẩn hóa mã: "Tân Đông" -> "AP_TAN_DONG"
        ap_code = ap_name.upper().replace(" ", "_")
        # Thay thế ký tự tiếng Việt
        for v, r in [("Ắ","A"),("Ằ","A"),("Ẵ","A"),("Ặ","A"),("Ẳ","A"),("Â","A"),("Ấ","A"),("Ầ","A"),("Ẫ","A"),("Ậ","A"),("Ẩ","A"),("Á","A"),("À","A"),("Ã","A"),("Ạ","A"),("Ả","A"),
                     ("É","E"),("È","E"),("Ẽ","E"),("Ẹ","E"),("Ẻ","E"),("Ê","E"),("Ế","E"),("Ề","E"),("Ễ","E"),("Ệ","E"),("Ể","E"),
                     ("Í","I"),("Ì","I"),("Ĩ","I"),("Ị","I"),("Ỉ","I"),
                     ("Ó","O"),("Ò","O"),("Õ","O"),("Ọ","O"),("Ỏ","O"),("Ô","O"),("Ố","O"),("Ồ","O"),("Ỗ","O"),("Ộ","O"),("Ổ","O"),("Ơ","O"),("Ớ","O"),("Ờ","O"),("Ỡ","O"),("Ợ","O"),("Ở","O"),
                     ("Ú","U"),("Ù","U"),("Ũ","U"),("Ụ","U"),("Ủ","U"),("Ư","U"),("Ứ","U"),("Ừ","U"),("Ữ","U"),("Ự","U"),("Ử","U"),
                     ("Ý","Y"),("Ỳ","Y"),("Ỹ","Y"),("Ỵ","Y"),("Ỷ","Y"),
                     ("Đ","D")]:
            ap_code = ap_code.replace(v, r)

        id_tuyen = f"AP_{ap_code}"

        # Kiểm tra đã tồn tại chưa
        cursor.execute("SELECT id_tuyen FROM tuyen_duong WHERE id_tuyen = ?", (id_tuyen,))
        if cursor.fetchone():
            return id_tuyen

        centroid = self.AP_CENTROIDS.get(ap_name, {"lat": 11.6050, "lng": 106.1850})
        import json
        cursor.execute("""
        INSERT OR IGNORE INTO tuyen_duong (
            id_tuyen, ten_duong, vi_tri, toado_batdau, toado_kethuc, phan_loai,
            diem_dau, diem_cuoi, chieu_dai_m, be_rong_m, ket_cau, tong_so_nha,
            quy_tac_danh_so, nguon_du_lieu
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            id_tuyen, f"Ấp {ap_name}", ap_name,
            json.dumps(centroid),
            json.dumps({"lat": centroid["lat"] + 0.003, "lng": centroid["lng"] + 0.003}),
            "Đường ấp (mặc định)", ap_name, ap_name,
            0.0, "3.0", "Nhựa / Đất / Hỗn hợp", 0,
            "Tả Chẵn - Hữu Lẻ", "Tự động tạo (nạp dữ liệu)"
        ))

        return id_tuyen

    def _tao_tuyen_duong_tinh_huyen(self, cursor, road_name: str, road_type: str) -> str:
        """Tạo tuyến đường cấp tỉnh/huyện (ĐT785, ĐH809...) nếu chưa có."""
        import json
        id_tuyen = f"TD_{road_name.replace(' ', '_').replace('.', '_')}"

        cursor.execute("SELECT id_tuyen FROM tuyen_duong WHERE id_tuyen = ?", (id_tuyen,))
        if cursor.fetchone():
            return id_tuyen

        phan_loai_map = {
            "ĐT": "Đường tỉnh", "DT": "Đường tỉnh",
            "ĐH": "Đường huyện", "DH": "Đường huyện",
            "QL": "Quốc lộ", "TL": "Tỉnh lộ"
        }
        phan_loai = phan_loai_map.get(road_type.upper(), "Đường tỉnh/huyện")

        centroid = {"lat": 11.6050, "lng": 106.1850}
        cursor.execute("""
        INSERT OR IGNORE INTO tuyen_duong (
            id_tuyen, ten_duong, vi_tri, toado_batdau, toado_kethuc, phan_loai,
            diem_dau, diem_cuoi, chieu_dai_m, be_rong_m, ket_cau, tong_so_nha,
            quy_tac_danh_so, nguon_du_lieu
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            id_tuyen, road_name, "Xã Tân Phú",
            json.dumps(centroid),
            json.dumps({"lat": centroid["lat"] + 0.01, "lng": centroid["lng"] + 0.01}),
            phan_loai, "Xã Tân Phú", "Xã Tân Phú",
            0.0, "6.0", "Nhựa", 0,
            "Tả Chẵn - Hữu Lẻ", "Tự động tạo (nạp dữ liệu)"
        ))

        return id_tuyen

    def _map_dia_chi_to_tuyen(self, cursor, dia_chi: str, ap: str) -> str:
        """
        Trích xuất tên đường từ địa chỉ hộ dân và map sang id_tuyen_duong.
        Thứ tự ưu tiên:
          1. Đường tỉnh/huyện: ĐT785, ĐH809, DT781B...
          2. Đường thu gom: THU35, THU92...
          3. Mặc định: tuyến cấp ấp
        """
        if not dia_chi:
            return self._tao_tuyen_duong_ap(cursor, ap)

        addr_upper = dia_chi.upper()

        # 1. Thử match đường tỉnh/huyện/quốc lộ
        m = self.ROAD_PATTERN.search(dia_chi)
        if m:
            road_type = m.group(1).upper()
            road_num = m.group(2).upper()
            road_name = f"{road_type}{road_num}"
            return self._tao_tuyen_duong_tinh_huyen(cursor, road_name, road_type)

        # 2. Thử match đường thu gom THU
        m_thu = self.THU_PATTERN.search(dia_chi)
        if m_thu:
            thu_num = m_thu.group(1)
            # Tìm id_tuyen trong DB với prefix THU.
            thu_code = thu_num.upper()
            # Các pattern trong DB: THU.35, THU.92, THU.88Đ...
            cursor.execute(
                "SELECT id_tuyen FROM tuyen_duong WHERE UPPER(ten_duong) LIKE ?",
                (f"THU.{thu_code}%",)
            )
            row = cursor.fetchone()
            if row:
                return row[0]
            # Không tìm thấy -> tạo mới
            return self._tao_tuyen_duong_tinh_huyen(cursor, f"THU.{thu_code}", "THU")

        # 3. Mặc định -> tuyến cấp ấp
        return self._tao_tuyen_duong_ap(cursor, ap)

    def _parse_thua_dat_lien_ket(self, lien_ket: Optional[str]) -> str:
        """
        Parse 'Tờ 106, Thửa 133 (CLN+ONT) | GCN: CN 447037'
        -> Trả về 'Tờ 106 - Thửa 133'
        """
        if not lien_ket:
            return ""
        m = re.match(r'Tờ\s*(\S+),?\s*Thửa\s*(\S+)', lien_ket)
        if m:
            return f"Tờ {m.group(1)} - Thửa {m.group(2)}"
        return lien_ket.split("|")[0].strip()

    def _parse_so_nha_cu(self, dia_chi: Optional[str]) -> Optional[str]:
        """Trích xuất số nhà cũ nếu địa chỉ có dạng 'Số nhà 191, DT785, ...'"""
        if not dia_chi:
            return None
        m = re.match(r'Số nhà\s+(\d+)', dia_chi, re.IGNORECASE)
        if m:
            return m.group(1)
        return None

    def nap_vao_so_nha(self) -> Dict[str, Any]:
        """
        Nạp 1.977 hộ dân đã đối soát (DU_DIEU_KIEN + CHUA_DU_DIEU_KIEN)
        vào bảng so_nha chính thức của hệ thống.

        Logic:
        - DU_DIEU_KIEN (1.281) → trang_thai='HIEN_HUU', muc_dich='DAT_O_THO_CU'
        - CHUA_DU_DIEU_KIEN (696) → trang_thai='NHA_CHUA_DU_DIEU_KIEN', muc_dich='DAT_CHUA_CHUYEN_MUC_DICH'
        - Map địa chỉ → id_tuyen_duong (ĐT/ĐH/THU → tuyến cụ thể, còn lại → tuyến ấp)
        - Parse thua_dat_lien_ket → so_to_thua
        - Generate GPS tọa độ theo ấp centroid + random offset
        """
        import json
        import random

        conn = self.storage.get_connection()
        cursor = conn.cursor()

        # Xóa dữ liệu so_nha cũ (nếu có) để nạp lại sạch
        cursor.execute("DELETE FROM so_nha")
        # Reset tong_so_nha trên tuyến đường
        cursor.execute("UPDATE tuyen_duong SET tong_so_nha = 0")

        # Lấy tất cả hộ dân đã đối soát
        cursor.execute("""
            SELECT id, stt_nha, stt_ho, ten_chu_ho, dia_chi, ap,
                   so_nhan_khau, ghi_chu, thua_dat_lien_ket, phan_loai_phap_ly
            FROM ho_dan
            WHERE phan_loai_phap_ly IN ('DU_DIEU_KIEN', 'CHUA_DU_DIEU_KIEN')
            ORDER BY ap, stt_nha, stt_ho
        """)
        all_ho = cursor.fetchall()

        total_inserted = 0
        du_dk_count = 0
        chua_du_dk_count = 0
        errors = []
        ap_counts = {}
        tuyen_counts = {}

        # Counter theo ấp để tạo id_nha và số nhà tạm
        ap_counter = {}

        for ho in all_ho:
            ho = dict(ho)
            ap = ho["ap"]
            phan_loai = ho["phan_loai_phap_ly"]

            # Counter per ấp
            if ap not in ap_counter:
                ap_counter[ap] = 0
            ap_counter[ap] += 1
            stt = ap_counter[ap]

            # === 1. Tạo id_nha ===
            ap_code = ap.upper().replace(" ", "")
            for v, r in [("Ắ","A"),("Ằ","A"),("Ẵ","A"),("Ặ","A"),("Ẳ","A"),("Â","A"),("Ấ","A"),("Ầ","A"),("Ẫ","A"),("Ậ","A"),("Ẩ","A"),("Á","A"),("À","A"),("Ã","A"),("Ạ","A"),("Ả","A"),
                         ("É","E"),("È","E"),("Ẽ","E"),("Ẹ","E"),("Ẻ","E"),("Ê","E"),("Ế","E"),("Ề","E"),("Ễ","E"),("Ệ","E"),("Ể","E"),
                         ("Í","I"),("Ì","I"),("Ĩ","I"),("Ị","I"),("Ỉ","I"),
                         ("Ó","O"),("Ò","O"),("Õ","O"),("Ọ","O"),("Ỏ","O"),("Ô","O"),("Ố","O"),("Ồ","O"),("Ỗ","O"),("Ộ","O"),("Ổ","O"),("Ơ","O"),("Ớ","O"),("Ờ","O"),("Ỡ","O"),("Ợ","O"),("Ở","O"),
                         ("Ú","U"),("Ù","U"),("Ũ","U"),("Ụ","U"),("Ủ","U"),("Ư","U"),("Ứ","U"),("Ừ","U"),("Ữ","U"),("Ự","U"),("Ử","U"),
                         ("Ý","Y"),("Ỳ","Y"),("Ỹ","Y"),("Ỵ","Y"),("Ỷ","Y"),
                         ("Đ","D")]:
                ap_code = ap_code.replace(v, r)
            id_nha = f"SN_{ap_code}_{stt:04d}"

            # === 2. Số nhà ===
            # Ưu tiên: số nhà cũ từ address > stt_nha từ Excel > counter
            so_nha_cu = self._parse_so_nha_cu(ho.get("dia_chi"))
            if so_nha_cu:
                so_nha_moi = f"{so_nha_cu}/{ap[:2]}"
            elif ho.get("stt_nha"):
                so_nha_moi = f"TMP-{ap[:2]}-{ho['stt_nha']}"
            else:
                so_nha_moi = f"TMP-{ap[:2]}-{stt}"

            # === 3. Tuyến đường ===
            id_tuyen = self._map_dia_chi_to_tuyen(cursor, ho.get("dia_chi", ""), ap)

            # === 4. Tọa độ GPS (centroid ấp + jitter) ===
            centroid = self.AP_CENTROIDS.get(ap, {"lat": 11.6050, "lng": 106.1850})
            random.seed(ho["id"])  # Deterministic per hộ
            lat = round(centroid["lat"] + random.uniform(-0.008, 0.008), 6)
            lng = round(centroid["lng"] + random.uniform(-0.008, 0.008), 6)
            toado_gps = json.dumps({"lat": lat, "lng": lng})

            # === 5. Trạng thái & Mục đích sử dụng đất ===
            if phan_loai == "DU_DIEU_KIEN":
                trang_thai = "HIEN_HUU"
                muc_dich = "DAT_O_THO_CU"
                hien_trang = "DA_XAY_NHA"
                du_dk_count += 1
            else:
                trang_thai = "NHA_CHUA_DU_DIEU_KIEN"
                muc_dich = "DAT_CHUA_CHUYEN_MUC_DICH"
                hien_trang = "DA_XAY_NHA"
                chua_du_dk_count += 1

            # === 6. Số tờ thửa ===
            so_to_thua = self._parse_thua_dat_lien_ket(ho.get("thua_dat_lien_ket"))

            # === 7. Ghi chú ===
            ghi_chu_parts = []
            if ho.get("ghi_chu"):
                ghi_chu_parts.append(ho["ghi_chu"])
            ghi_chu_parts.append(f"Ấp {ap}")
            if ho.get("dia_chi"):
                ghi_chu_parts.append(f"ĐC: {ho['dia_chi']}")
            if ho.get("thua_dat_lien_ket"):
                ghi_chu_parts.append(f"Đất: {ho['thua_dat_lien_ket']}")
            ghi_chu = " | ".join(ghi_chu_parts)

            try:
                cursor.execute("""
                INSERT INTO so_nha (
                    id_nha, so_nha_moi, so_nha_cu, ten_chu_ho, cccd, sdt,
                    toado_gps, trang_thai, muc_dich_su_dung_dat, hien_trang_xay_dung,
                    so_to_thua, id_tuyen_duong, id_hem_ngach, loai_nha,
                    dien_tich_mat_tien, ngay_cap, so_giay_phep, ghi_chu
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (
                    id_nha, so_nha_moi, so_nha_cu, ho["ten_chu_ho"],
                    None, None,  # cccd, sdt - chưa có
                    toado_gps, trang_thai, muc_dich, hien_trang,
                    so_to_thua, id_tuyen, None,  # id_hem_ngach
                    "Nhà ở riêng lẻ", 5.0,
                    None, None,  # ngay_cap, so_giay_phep
                    ghi_chu
                ))

                # Cập nhật tong_so_nha trên tuyến
                cursor.execute(
                    "UPDATE tuyen_duong SET tong_so_nha = tong_so_nha + 1 WHERE id_tuyen = ?",
                    (id_tuyen,)
                )

                total_inserted += 1
                ap_counts[ap] = ap_counts.get(ap, 0) + 1
                tuyen_counts[id_tuyen] = tuyen_counts.get(id_tuyen, 0) + 1

            except Exception as e:
                errors.append(f"{ho['ten_chu_ho']} ({ap}): {str(e)}")

        # Log kết quả
        cursor.execute(
            "INSERT INTO log_hethong (hanh_dong, doi_tuong, chi_tiet) VALUES (?, ?, ?)",
            ("NẠP DỮ LIỆU VÀO SỐ NHÀ", "so_nha",
             f"Đã nạp {total_inserted} bản ghi ({du_dk_count} đủ ĐK, {chua_du_dk_count} chưa đủ ĐK). "
             f"Lỗi: {len(errors)}. "
             f"Tuyến đường liên quan: {len(tuyen_counts)}")
        )
        conn.commit()
        conn.close()

        self.storage._cache.clear()

        return {
            "success": True,
            "message": (
                f"✅ Đã nạp thành công {total_inserted} bản ghi vào hệ thống số nhà!\n"
                f"   • Đủ điều kiện cấp số: {du_dk_count}\n"
                f"   • Chưa đủ điều kiện (ghi nhận): {chua_du_dk_count}\n"
                f"   • Số tuyến đường liên quan: {len(tuyen_counts)}\n"
                f"   • Lỗi: {len(errors)}"
            ),
            "total_inserted": total_inserted,
            "du_dieu_kien": du_dk_count,
            "chua_du_dieu_kien": chua_du_dk_count,
            "theo_ap": ap_counts,
            "so_tuyen_duong": len(tuyen_counts),
            "top_tuyen": dict(sorted(tuyen_counts.items(), key=lambda x: -x[1])[:15]),
            "errors": errors[:10] if errors else []
        }

    # ====================== NẠP ĐẤT THỔ CƯ CHƯA XÂY NHÀ ======================

    def nap_dat_tho_cu_chua_xay(self) -> Dict[str, Any]:
        """
        Nạp ~1.308 thửa đất thổ cư (co_dat_o=1) CHƯA được liên kết với hộ dân nào
        vào bảng so_nha với trạng thái DAT_THO_CU_CHUA_XAY.

        Mục đích: Dự trữ quỹ số nhà cho các thửa đất ở chưa xây dựng,
        phục vụ quy hoạch và cấp số nhà tương lai khi có nhà mới xây.

        Logic:
        - Lấy tất cả thửa đất có co_dat_o=1 mà KHÔNG khớp với bất kỳ hộ dân nào
          (dựa trên cột thua_dat_lien_ket của ho_dan)
        - Tạo bản ghi so_nha tạm với id_nha = DT_{so_to}_{so_thua}
        - Không xóa dữ liệu so_nha cũ (append only)
        """
        conn = self.storage.get_connection()
        cursor = conn.cursor()

        # Lấy các thửa đất thổ cư KHÔNG liên kết với hộ dân nào
        # Bảng ho_dan.thua_dat_lien_ket chứa dạng: "Tờ 106, Thửa 133 (CLN+ONT) | GCN: ..."
        # Nên ta dùng LIKE '%Tờ {so_to}, Thửa {so_thua}%' để đối soát
        cursor.execute("""
            SELECT td.id, td.so_to, td.so_thua, td.ten_chu, td.ma_loai_dat,
                   td.dien_tich, td.so_gcn
            FROM thua_dat td
            WHERE td.co_dat_o = 1
              AND td.id NOT IN (
                  SELECT td2.id
                  FROM thua_dat td2
                  JOIN ho_dan hd ON hd.thua_dat_lien_ket LIKE '%Tờ ' || td2.so_to || ', Thửa ' || td2.so_thua || '%'
                  WHERE td2.co_dat_o = 1
              )
            ORDER BY CAST(td.so_to AS INTEGER), CAST(td.so_thua AS INTEGER)
        """)
        parcels = cursor.fetchall()

        # Tạo tuyến đường mặc định cho đất chưa xác định ấp
        id_tuyen_mac_dinh = "AP_CHUA_XAC_DINH"
        cursor.execute("SELECT id_tuyen FROM tuyen_duong WHERE id_tuyen = ?", (id_tuyen_mac_dinh,))
        if not cursor.fetchone():
            centroid_default = {"lat": 11.6050, "lng": 106.1850}
            cursor.execute("""
            INSERT OR IGNORE INTO tuyen_duong (
                id_tuyen, ten_duong, vi_tri, toado_batdau, toado_kethuc, phan_loai,
                diem_dau, diem_cuoi, chieu_dai_m, be_rong_m, ket_cau, tong_so_nha,
                quy_tac_danh_so, nguon_du_lieu
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                id_tuyen_mac_dinh, "Chưa xác định ấp", "Xã Tân Phú",
                json.dumps(centroid_default),
                json.dumps({"lat": centroid_default["lat"] + 0.003, "lng": centroid_default["lng"] + 0.003}),
                "Đất thổ cư chưa xây (dự trữ)", "Xã Tân Phú", "Xã Tân Phú",
                0.0, "0", "Chưa xác định", 0,
                "Chưa áp dụng", "Tự động tạo (nạp đất thổ cư chưa xây)"
            ))

        total_inserted = 0
        errors = []
        # Centroid mặc định xã Tân Phú để jitter tọa độ
        centroid_xa = {"lat": 11.6050, "lng": 106.1850}

        for parcel in parcels:
            p = dict(parcel)
            so_to = str(p["so_to"]).strip()
            so_thua = str(p["so_thua"]).strip()
            thua_dat_id = p["id"]

            # === 1. Tạo id_nha và số nhà tạm ===
            id_nha = f"DT_{so_to}_{so_thua}"
            so_nha_moi = f"DT-{so_to}/{so_thua}"

            # === 2. Tên chủ hộ ===
            ten_chu_ho = normalize_name(p.get("ten_chu")) or "Chưa xác định"

            # === 3. Trạng thái và phân loại ===
            trang_thai = "DAT_THO_CU_CHUA_XAY"
            muc_dich = "DAT_O_THO_CU"
            hien_trang = "DAT_TRONG"

            # === 4. Số tờ thửa ===
            so_to_thua = f"Tờ {so_to} - Thửa {so_thua}"

            # === 5. Tọa độ GPS (centroid xã + jitter deterministic) ===
            random.seed(thua_dat_id)
            lat = round(centroid_xa["lat"] + random.uniform(-0.008, 0.008), 6)
            lng = round(centroid_xa["lng"] + random.uniform(-0.008, 0.008), 6)
            toado_gps = json.dumps({"lat": lat, "lng": lng})

            # === 6. Ghi chú chi tiết ===
            ma_loai_dat = p.get("ma_loai_dat", "")
            dien_tich = p.get("dien_tich", 0)
            so_gcn = p.get("so_gcn") or "Không"
            ghi_chu = (
                f"Đất thổ cư chưa xây nhà | "
                f"Loại đất: {ma_loai_dat} | "
                f"Diện tích: {dien_tich}m² | "
                f"GCN: {so_gcn}"
            )

            try:
                cursor.execute("""
                INSERT OR IGNORE INTO so_nha (
                    id_nha, so_nha_moi, so_nha_cu, ten_chu_ho, cccd, sdt,
                    toado_gps, trang_thai, muc_dich_su_dung_dat, hien_trang_xay_dung,
                    so_to_thua, id_tuyen_duong, id_hem_ngach, loai_nha,
                    dien_tich_mat_tien, ngay_cap, so_giay_phep, ghi_chu
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (
                    id_nha, so_nha_moi, None, ten_chu_ho,
                    None, None,  # cccd, sdt - chưa có
                    toado_gps, trang_thai, muc_dich, hien_trang,
                    so_to_thua, id_tuyen_mac_dinh, None,  # id_hem_ngach
                    None, 0.0,  # loai_nha, dien_tich_mat_tien
                    None, None,  # ngay_cap, so_giay_phep
                    ghi_chu
                ))

                if cursor.rowcount > 0:
                    total_inserted += 1
                    # Cập nhật tong_so_nha trên tuyến
                    cursor.execute(
                        "UPDATE tuyen_duong SET tong_so_nha = tong_so_nha + 1 WHERE id_tuyen = ?",
                        (id_tuyen_mac_dinh,)
                    )

            except Exception as e:
                errors.append(f"Tờ {so_to} Thửa {so_thua}: {str(e)}")

        # Log kết quả vào log_hethong
        cursor.execute(
            "INSERT INTO log_hethong (hanh_dong, doi_tuong, chi_tiet) VALUES (?, ?, ?)",
            ("NẠP ĐẤT THỔ CƯ CHƯA XÂY", "so_nha",
             f"Đã nạp {total_inserted}/{len(parcels)} thửa đất thổ cư chưa xây nhà vào so_nha. "
             f"Lỗi: {len(errors)}")
        )
        conn.commit()
        conn.close()

        self.storage._cache.clear()

        return {
            "success": True,
            "message": (
                f"✅ Đã nạp {total_inserted} thửa đất thổ cư chưa xây nhà vào hệ thống!\n"
                f"   • Tổng thửa đất thổ cư không liên kết hộ dân: {len(parcels)}\n"
                f"   • Đã thêm mới: {total_inserted}\n"
                f"   • Tuyến mặc định: {id_tuyen_mac_dinh}\n"
                f"   • Lỗi: {len(errors)}"
            ),
            "total_parcels_found": len(parcels),
            "total_inserted": total_inserted,
            "errors": errors[:10] if errors else []
        }
