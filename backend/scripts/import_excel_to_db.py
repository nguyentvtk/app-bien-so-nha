#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
IMPORT EXCEL RÀ SOÁT ĐẤT ĐAI / HỘ DÂN -> SQLITE (sonha_database.db)
===================================================================

Trích xuất toàn bộ dữ liệu từ các file .xlsx trong thư mục dự án vào bảng
`ho_dan_dat_dai` của SQLite, gồm các trường chính:
    ID, Họ tên chủ hộ, SĐT, Địa chỉ hiện trạng, Số tờ, Số thửa,
    Diện tích, Loại đất, Tình trạng GCN, Tọa độ GPS (nếu có).

Script TỰ NHẬN DIỆN 2 loại cấu trúc file nguồn (theo nội dung, không phụ thuộc tên file):

1. File bảng kê hộ dân (VD: DS_Nha_thuc_te.xlsx)
   - Mỗi sheet = 1 ấp; dòng tiêu đề: "STT nhà | STT hộ | Họ và tên chủ hộ |
     Địa chỉ của hộ | Số nhân khẩu của hộ khi lập bảng kê | Ghi chú".
   - Sheet danh sách công chức/tổ trưởng (không phải hộ dân) được bỏ qua.
   - Địa chỉ ấp được lấy từ dòng "Ấp/Khu phố: ..." trong đầu sheet.

2. File thống kê rà soát thửa đất theo Kế hoạch 2959/KH-BNNMT-BCA
   (VD: Trang_1.xlsx, Trang_2.xlsx, Trang_3.xlsx) — tiêu đề 3 tầng (R4-R6),
   dữ liệu từ dòng 7:
     B  Tên đơn vị hành chính cấp xã      P  Thông tin chủ (chủ đất)
     C  Số tờ bản đồ                      Q  Thông tin giấy chứng nhận
     D  Số thứ tự thửa đất                R  Đã định danh thửa
     E-N Tiêu chí đánh giá (dấu X)        S  Đã quy đổi
     O  Phân loại dữ liệu (Nhóm 1/2/3)    T  Thông tin thiếu
     U  Diện tích                         V  Mã loại đất (ONT, HNK, CLN...)

Ghi chú về nguồn dữ liệu hiện tại:
    - KHÔNG file nào chứa tọa độ GPS  -> cột `toa_do_gps` giữ NULL (nếu có sẽ nhập).
    - KHÔNG có SĐT của từng hộ dân    -> cột `so_dien_thoai` NULL
      (script vẫn tự dò SĐT trong cột Ghi chú phòng khi dữ liệu bổ sung sau).
    - Tình trạng GCN được suy ra từ cột Q và cột T:
        * "không có giấy chứng nhận"              -> "Chưa có GCN"
        * cột Q "CG 954942" hoặc T "có giấy chứng nhận 4564434_0"
                                                 -> "Đã cấp GCN (số ...)"
        * các cụm "không có ..." trong T          -> gộp thành "Thiếu: ..."

Sử dụng:
    python3 backend/scripts/import_excel_to_db.py                # import mặc định
    python3 backend/scripts/import_excel_to_db.py --reset        # xoá bảng, nhập lại từ đầu
    python3 backend/scripts/import_excel_to_db.py --no-enrich    # tắt ghép địa chỉ cho thửa đất
    python3 backend/scripts/import_excel_to_db.py --excel-dir DIR --db PATH [--verbose]

Chạy idempotent: các dòng đã nhập (cùng file/sheet/dòng nguồn) sẽ bị bỏ qua
 nhờ UNIQUE INDEX, trừ khi dùng --reset.
"""

from __future__ import annotations

import argparse
import re
import sqlite3
import sys
from pathlib import Path

import openpyxl

# ---------------------------------------------------------------------------
# Đường dẫn mặc định (tính tương đối so với vị trí script, chạy được mọi nơi)
#   backend/scripts/import_excel_to_db.py  ->  backend/data/sonha_database.db
#                                            và thư mục gốc dự án (chứa .xlsx)
# ---------------------------------------------------------------------------
SCRIPT_DIR = Path(__file__).resolve().parent
BACKEND_DIR = SCRIPT_DIR.parent
PROJECT_DIR = BACKEND_DIR.parent
DEFAULT_DB = BACKEND_DIR / "data" / "sonha_database.db"
DEFAULT_EXCEL_DIR = PROJECT_DIR

PHONE_RE = re.compile(r"(?:0|\+84)[\s.\-]?\d{8,10}")
AP_RE = re.compile(r"Ấp\s*/?\s*Khu phố\s*:?\s*(.+)", re.IGNORECASE)
GCN_SO_RE = re.compile(r"có giấy chứng nhận\s+([A-Za-z0-9_/\-.]+)", re.IGNORECASE)
THIEU_RE = re.compile(r"không có\s+([^;]+)", re.IGNORECASE)
TONG_CONG_RE = re.compile(r"^\s*(tổng|cộng)", re.IGNORECASE)


# ===========================================================================
# Hàm tiện ích làm sạch dữ liệu
# ===========================================================================
def clean_text(value) -> str | None:
    """Chuyển giá trị ô Excel thành chuỗi sạch: bỏ xuống dòng, rút gọn khoảng trắng."""
    if value is None:
        return None
    text = str(value).replace("\n", " ").strip()
    text = re.sub(r"\s{2,}", " ", text)
    return text or None


def to_float(value) -> float | None:
    if value is None:
        return None
    try:
        return float(str(value).replace(",", "."))
    except ValueError:
        return None


def to_int(value) -> int | None:
    if value is None:
        return None
    try:
        return int(float(str(value)))
    except ValueError:
        return None


def cell(row: tuple, idx: int):
    """Lấy ô an toàn theo chỉ số cột (dòng read_only có thể ngắn hơn số cột)."""
    return row[idx] if idx < len(row) else None


def norm_name(name: str | None) -> str | None:
    """Chuẩn hoá tên riêng để so khớp: thường hoá, bỏ "Bà/Ông", gọn khoảng trắng."""
    if not name:
        return None
    text = re.sub(r"\s{2,}", " ", name.strip().lower())
    text = re.sub(r"^(bà|ông|anh|chị|cô|mr|mrs)\s+", "", text)
    return text.strip(" .,;")


def derive_gcn(thong_tin_gcn: str | None, thong_tin_thieu: str | None) -> str | None:
    """Suy ra tình trạng GCN từ cột Q (thông tin GCN) và cột T (thông tin thiếu)."""
    if thong_tin_gcn and thong_tin_gcn.strip():
        so = thong_tin_gcn.replace("CG", "").strip()
        base = f"Đã cấp GCN (số {so})" if so else f"Đã cấp GCN ({thong_tin_gcn.strip()})"
    elif thong_tin_thieu and GCN_SO_RE.search(thong_tin_thieu):
        base = f"Đã cấp GCN (số {GCN_SO_RE.search(thong_tin_thieu).group(1)})"
    elif thong_tin_thieu and "không có giấy chứng nhận" in thong_tin_thieu.lower():
        return "Chưa có GCN"
    else:
        return None

    # Gộp các thiếu sót ("không có ngày vào sổ", "không có hồ sơ quét"...)
    # thành phần đuôi "Thiếu: ..." để cán bộ dễ đọc.
    if thong_tin_thieu:
        gaps = []
        for match in THIEU_RE.finditer(thong_tin_thieu):
            gap = match.group(1).strip().rstrip(".")
            # Cắt phần lồng nhau kiểu "dữ liệu không gian, không có hồ sơ quét"
            gap = re.split(r",\s*không có", gap, flags=re.IGNORECASE)[0]
            if gap and gap not in gaps and "giấy chứng nhận" not in gap.lower():
                gaps.append(gap)
        if gaps:
            return base + " — Thiếu: " + "; ".join(gaps[:4])
    return base


# ===========================================================================
# Phát hiện cấu trúc & parse từng loại sheet
# ===========================================================================
def detect_sheet_kind(ws) -> str | None:
    """Dò 10 dòng đầu để nhận dạng: 'HO_DAN' | 'THUA_DAT' | None."""
    for row in ws.iter_rows(min_row=1, max_row=10, values_only=True):
        texts = [clean_text(v) or "" for v in row]
        joined = " | ".join(texts)
        if "Họ và tên chủ hộ" in joined and "Địa chỉ của hộ" in joined:
            return "HO_DAN"
        if "BẢNG THỐNG KÊ KẾT QUẢ RÀ SOÁT" in joined:
            return "THUA_DAT"
        if "Số tờ bản đồ" in joined and "Thông tin chủ" in joined:
            return "THUA_DAT"
    return None


def parse_ho_dan_sheet(ws, file_name: str) -> list[dict]:
    """Parse sheet bảng kê hộ dân của 1 ấp."""
    header_row = None
    ap_name = None
    rows_cache: list[tuple] = []
    for r_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
        rows_cache.append(row)
        texts = [clean_text(v) or "" for v in row]
        joined = " | ".join(texts)
        if header_row is None and "Họ và tên chủ hộ" in joined:
            header_row = r_idx
        if ap_name is None:
            for text in texts:
                if text and (m := AP_RE.search(text)):
                    ap_name = clean_text(m.group(1))
    if header_row is None:
        return []
    if ap_name is None:
        # Dự phòng: lấy tên ấp từ tên sheet dạng "1,Tân Đông"
        ap_name = clean_text(ws.title.split(",", 1)[-1])

    records = []
    for r_idx, row in enumerate(rows_cache, 1):
        if r_idx <= header_row:
            continue
        ho_ten = clean_text(cell(row, 2))
        if not ho_ten:
            continue
        # Bỏ dòng mã cột (A|B|C|D) và dòng tổng cộng cuối bảng
        if clean_text(cell(row, 0)) in {"A", "B", "C", "D"} and clean_text(cell(row, 1)) in {"A", "B", "C", "D"}:
            continue
        if TONG_CONG_RE.match(ho_ten):
            continue
        ghi_chu = clean_text(cell(row, 5))
        records.append({
            "loai_ban_ghi": "HO_DAN",
            "ho_ten_chu_ho": ho_ten,
            "so_dien_thoai": (PHONE_RE.search(ghi_chu).group(0) if ghi_chu and PHONE_RE.search(ghi_chu) else None),
            "dia_chi_hien_trang": clean_text(cell(row, 3)),
            "so_to": None,
            "so_thua": None,
            "dien_tich": None,
            "loai_dat": None,
            "tinh_trang_gcn": None,
            "toa_do_gps": None,
            "xa": "Tân Phú",
            "ap": ap_name,
            "stt_nha": to_int(cell(row, 0)),
            "stt_ho": to_int(cell(row, 1)),
            "so_nhan_khau": to_int(cell(row, 4)),
            "phan_loai_du_lieu": None,
            "thong_tin_thieu": None,
            "ghi_chu": ghi_chu,
            "nguon_file": file_name,
            "sheet_nguon": clean_text(ws.title),
            "dong_nguon": r_idx,
        })
    return records


def parse_thua_dat_sheet(ws, file_name: str) -> list[dict]:
    """Parse sheet thống kê rà soát thửa đất (KH 2959) — header 3 tầng, cột theo tên lá.

    Bảng có 3 dòng tiêu đề ghép ô (nhóm / nhóm con / lá); nhãn "lá" của từng cột
    nằm ở dòng khác nhau (C-N ở dòng dưới, P-V ở dòng giữa) nên phải quét cả 3
    dòng và cho dòng sâu nhất (lá) ghi đè."""
    rows_cache: list[tuple] = [(None,)]
    leaf_header_row = None
    for r_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
        rows_cache.append(row)
        texts = [clean_text(v) or "" for v in row]
        if leaf_header_row is None and any(t == "Số tờ bản đồ" for t in texts):
            leaf_header_row = r_idx

    if leaf_header_row is None:
        return []

    # Ánh xạ tên cột -> chỉ số, quét từ dòng 1 tới CẢ dòng lá; ghi đè để nhãn lá thắng
    col_map: dict[int, str] = {}
    for row in rows_cache[1:leaf_header_row + 1]:
        for c_idx, v in enumerate(row):
            text = clean_text(v)
            if text:
                col_map[c_idx] = text

    def col_of(*prefixes: str) -> int | None:
        for prefix in prefixes:
            for c_idx, name in col_map.items():
                if name == prefix or name.startswith(prefix):
                    return c_idx
        return None

    c_stt = col_of("STT")
    c_xa = col_of("Tên đơn vị hành chính cấp xã")
    c_to = col_of("Số tờ bản đồ")
    c_thua = col_of("Số thứ tự thửa đất")
    c_phanloai = col_of("Phân loại dữ liệu")
    c_chu = col_of("Thông tin chủ")
    c_gcn = col_of("Thông tin giấy chứng nhận")
    c_dinh_danh = col_of("Đã định danh thửa")
    c_quy_doi = col_of("Đã quy đổi")
    c_thieu = col_of("Thông tin thiếu")
    c_dt = col_of("Diện tích")
    c_loai = col_of("Mã loại đất")

    # Các cột tiêu chí đánh giá (E..N): mọi cột còn lại không thuộc nhóm trên
    known_cols = {c_stt, c_xa, c_to, c_thua, c_phanloai, c_chu, c_gcn,
                  c_dinh_danh, c_quy_doi, c_thieu, c_dt, c_loai}
    tieu_chi_cols = {
        idx: name for idx, name in col_map.items()
        if idx not in known_cols and "Tiêu chí" not in name and "Thông tin bổ sung" not in name
    }

    records = []
    for r_idx, row in enumerate(rows_cache, 1):
        if r_idx <= leaf_header_row:
            continue
        if c_chu is not None:
            chu = clean_text(cell(row, c_chu))
        else:
            chu = None
        stt = clean_text(cell(row, c_stt)) if c_stt is not None else None
        if not chu and not stt:
            continue
        if TONG_CONG_RE.match(chu or "") or (stt and not stt.isdigit() and not chu):
            continue

        thong_tin_thieu = clean_text(cell(row, c_thieu)) if c_thieu is not None else None
        thong_tin_gcn = clean_text(cell(row, c_gcn)) if c_gcn is not None else None

        # Gom các tiêu chí đánh giá được đánh dấu "X" thành ghi chú đọc được
        danh_dau = []
        for idx, ten in tieu_chi_cols.items():
            if (clean_text(cell(row, idx)) or "").upper() == "X":
                danh_dau.append(ten)
        ghi_chu_parts = [f"Tiêu chí đạt: {'; '.join(danh_dau)}"] if danh_dau else []
        if c_dinh_danh is not None and (clean_text(cell(row, c_dinh_danh)) or "").upper() == "X":
            ghi_chu_parts.append("Đã định danh thửa")
        if c_quy_doi is not None and (clean_text(cell(row, c_quy_doi)) or "").upper() == "X":
            ghi_chu_parts.append("Đã quy đổi")

        records.append({
            "loai_ban_ghi": "THUA_DAT",
            "ho_ten_chu_ho": chu,
            "so_dien_thoai": None,
            "dia_chi_hien_trang": None,  # sẽ được làm giàu từ danh sách hộ dân
            "so_to": clean_text(cell(row, c_to)) if c_to is not None else None,
            "so_thua": clean_text(cell(row, c_thua)) if c_thua is not None else None,
            "dien_tich": to_float(cell(row, c_dt)) if c_dt is not None else None,
            "loai_dat": clean_text(cell(row, c_loai)) if c_loai is not None else None,
            "tinh_trang_gcn": derive_gcn(thong_tin_gcn, thong_tin_thieu),
            "toa_do_gps": None,
            "xa": clean_text(cell(row, c_xa)) if c_xa is not None else None,
            "ap": None,
            "stt_nha": None,
            "stt_ho": to_int(stt),
            "so_nhan_khau": None,
            "phan_loai_du_lieu": clean_text(cell(row, c_phanloai)) if c_phanloai is not None else None,
            "thong_tin_thieu": thong_tin_thieu,
            "ghi_chu": " | ".join(ghi_chu_parts) or None,
            "nguon_file": file_name,
            "sheet_nguon": clean_text(ws.title),
            "dong_nguon": r_idx,
        })
    return records


# ===========================================================================
# Cơ sở dữ liệu
# ===========================================================================
DDL = """
CREATE TABLE IF NOT EXISTS ho_dan_dat_dai (
    id                 INTEGER PRIMARY KEY AUTOINCREMENT,
    ho_ten_chu_ho      TEXT,
    so_dien_thoai      TEXT,
    dia_chi_hien_trang TEXT,
    so_to              TEXT,
    so_thua            TEXT,
    dien_tich          REAL,
    loai_dat           TEXT,
    tinh_trang_gcn     TEXT,
    toa_do_gps         TEXT,
    loai_ban_ghi       TEXT NOT NULL CHECK (loai_ban_ghi IN ('HO_DAN', 'THUA_DAT')),
    xa                 TEXT,
    ap                 TEXT,
    stt_nha            INTEGER,
    stt_ho             INTEGER,
    so_nhan_khau       INTEGER,
    phan_loai_du_lieu  TEXT,
    thong_tin_thieu    TEXT,
    ghi_chu            TEXT,
    nguon_file         TEXT NOT NULL,
    sheet_nguon        TEXT,
    dong_nguon         INTEGER NOT NULL
);
"""

INSERT_SQL = """
INSERT OR IGNORE INTO ho_dan_dat_dai (
    ho_ten_chu_ho, so_dien_thoai, dia_chi_hien_trang, so_to, so_thua,
    dien_tich, loai_dat, tinh_trang_gcn, toa_do_gps, loai_ban_ghi,
    xa, ap, stt_nha, stt_ho, so_nhan_khau, phan_loai_du_lieu,
    thong_tin_thieu, ghi_chu, nguon_file, sheet_nguon, dong_nguon
) VALUES (
    :ho_ten_chu_ho, :so_dien_thoai, :dia_chi_hien_trang, :so_to, :so_thua,
    :dien_tich, :loai_dat, :tinh_trang_gcn, :toa_do_gps, :loai_ban_ghi,
    :xa, :ap, :stt_nha, :stt_ho, :so_nhan_khau, :phan_loai_du_lieu,
    :thong_tin_thieu, :ghi_chu, :nguon_file, :sheet_nguon, :dong_nguon
);
"""


def enrich_dia_chi_thua_dat(conn: sqlite3.Connection) -> int:
    """Làm giàu địa chỉ cho bản ghi thửa đất: khớp TUYỆT ĐỐI theo tên chủ
    (đã chuẩn hoá, bỏ "Bà/Ông") với đúng 1 hộ dân duy nhất -> mới gán địa chỉ.
    Trả về số bản ghi được bổ sung địa chỉ."""
    ho_dan = {}
    for row_id, ho_ten, dia_chi, ap in conn.execute(
        "SELECT id, ho_ten_chu_ho, dia_chi_hien_trang, ap FROM ho_dan_dat_dai "
        "WHERE loai_ban_ghi = 'HO_DAN'"
    ):
        key = norm_name(ho_ten)
        if key:
            # Chỉ giữ tên khớp duy nhất; trùng nhiều hộ -> loại khỏi đồ khớp
            if key in ho_dan:
                ho_dan[key] = None
            else:
                ho_dan[key] = (row_id, dia_chi, ap)

    matched = 0
    for row_id, ho_ten in conn.execute(
        "SELECT id, ho_ten_chu_ho FROM ho_dan_dat_dai "
        "WHERE loai_ban_ghi = 'THUA_DAT' AND (dia_chi_hien_trang IS NULL OR dia_chi_hien_trang = '')"
    ).fetchall():
        info = ho_dan.get(norm_name(ho_ten))
        if not info:
            continue
        _, dia_chi, ap = info
        conn.execute(
            "UPDATE ho_dan_dat_dai SET dia_chi_hien_trang = COALESCE(dia_chi_hien_trang, ?), "
            "ap = COALESCE(ap, ?) WHERE id = ?",
            (dia_chi, ap, row_id),
        )
        matched += 1
    conn.commit()
    return matched


# ===========================================================================
# Main
# ===========================================================================
def main() -> int:
    parser = argparse.ArgumentParser(
        description="Nhập dữ liệu Excel rà soát đất đai / hộ dân vào SQLite."
    )
    parser.add_argument("--excel-dir", type=Path, default=DEFAULT_EXCEL_DIR,
                        help=f"Thư mục chứa file .xlsx (mặc định: {DEFAULT_EXCEL_DIR})")
    parser.add_argument("--db", type=Path, default=DEFAULT_DB,
                        help=f"Đường dẫn SQLite (mặc định: {DEFAULT_DB})")
    parser.add_argument("--reset", action="store_true",
                        help="Xoá bảng cũ và nhập lại toàn bộ từ đầu")
    parser.add_argument("--no-enrich", action="store_true",
                        help="Tắt bước ghép địa chỉ hộ dân cho bản ghi thửa đất")
    parser.add_argument("--verbose", action="store_true",
                        help="In chi tiết từng sheet")
    args = parser.parse_args()

    excel_files = sorted(
        f for f in args.excel_dir.glob("*.xlsx")
        if not f.name.startswith("~$") and not f.name.startswith(".")
    )
    if not excel_files:
        print(f"[LỖI] Không tìm thấy file .xlsx nào trong: {args.excel_dir}")
        return 1

    args.db.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(args.db)
    if args.reset:
        conn.execute("DROP INDEX IF EXISTS ux_nguon_dong")
        conn.execute("DROP INDEX IF EXISTS ix_ho_ten")
        conn.execute("DROP TABLE IF EXISTS ho_dan_dat_dai")
        print("[RESET] Đã xoá bảng ho_dan_dat_dai cũ.")
    conn.execute(DDL)
    conn.execute("CREATE UNIQUE INDEX IF NOT EXISTS ux_nguon_dong "
                 "ON ho_dan_dat_dai (nguon_file, sheet_nguon, dong_nguon)")
    conn.execute("CREATE INDEX IF NOT EXISTS ix_ho_ten ON ho_dan_dat_dai (ho_ten_chu_ho)")
    conn.execute("CREATE INDEX IF NOT EXISTS ix_ap ON ho_dan_dat_dai (ap)")

    total = {"HO_DAN": 0, "THUA_DAT": 0}
    skipped_sheets: list[str] = []

    for excel_file in excel_files:
        print(f"\n=== File: {excel_file.name} ===")
        wb = openpyxl.load_workbook(excel_file, read_only=True, data_only=True)
        for ws in wb.worksheets:
            kind = detect_sheet_kind(ws)
            if kind is None:
                skipped_sheets.append(f"{excel_file.name} :: {ws.title}")
                continue
            records = (parse_ho_dan_sheet if kind == "HO_DAN" else parse_thua_dat_sheet)(
                ws, excel_file.name
            )
            before = conn.total_changes
            conn.executemany(INSERT_SQL, records)
            inserted = conn.total_changes - before
            total[kind] += inserted
            label = "hộ dân" if kind == "HO_DAN" else "thửa đất"
            ap_info = f", ấp: {records[0]['ap']}" if kind == "HO_DAN" and records and records[0]["ap"] else ""
            print(f"  [{label:>9}] sheet '{ws.title}': đọc {len(records)} dòng, "
                  f"nhập mới {inserted}{ap_info}")
            if args.verbose:
                for rec in records[:3]:
                    print(f"      • {rec['ho_ten_chu_ho']} | {rec['dia_chi_hien_trang'] or ''} | "
                          f"tờ {rec['so_to'] or '-'} thửa {rec['so_thua'] or '-'}")
        wb.close()

    conn.commit()

    enriched = 0
    if not args.no_enrich and total["THUA_DAT"] >= 0:
        enriched = enrich_dia_chi_thua_dat(conn)

    # ------------------------------------------------------------------
    # Thống kê tổng kết
    # ------------------------------------------------------------------
    n_all, n_ho, n_thua = conn.execute(
        "SELECT COUNT(*), "
        "SUM(loai_ban_ghi='HO_DAN'), SUM(loai_ban_ghi='THUA_DAT') "
        "FROM ho_dan_dat_dai"
    ).fetchone()
    n_co_gcn = conn.execute(
        "SELECT COUNT(*) FROM ho_dan_dat_dai WHERE tinh_trang_gcn LIKE 'Đã cấp%'"
    ).fetchone()[0]
    n_chua_gcn = conn.execute(
        "SELECT COUNT(*) FROM ho_dan_dat_dai WHERE tinh_trang_gcn LIKE 'Chưa có%'"
    ).fetchone()[0]
    ap_stats = conn.execute(
        "SELECT ap, COUNT(*) FROM ho_dan_dat_dai "
        "WHERE loai_ban_ghi='HO_DAN' AND ap IS NOT NULL GROUP BY ap ORDER BY ap"
    ).fetchall()

    print("\n" + "=" * 62)
    print(f"HOÀN TẤT — CSDL: {args.db}")
    print(f"  Tổng số bản ghi        : {n_all}")
    print(f"  ├── Hộ dân             : {n_ho}")
    print(f"  └── Thửa đất rà soát   : {n_thua} (GCN đã cấp: {n_co_gcn}, chưa có: {n_chua_gcn})")
    if not args.no_enrich:
        print(f"  Thửa đất được ghép địa chỉ theo tên chủ: {enriched}")
    if ap_stats:
        print("  Số hộ dân theo ấp:")
        for ap, count in ap_stats:
            print(f"    - {ap:<14}: {count}")
    if skipped_sheets:
        print("  Sheet bỏ qua (không phải dữ liệu hộ/thửa):")
        for s in skipped_sheets:
            print(f"    - {s}")
    print("  Lưu ý: nguồn Excel hiện không có tọa độ GPS và SĐT hộ dân"
          " → 2 cột này đang NULL, sẵn sàng nhận khi bổ sung dữ liệu.")
    conn.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())
